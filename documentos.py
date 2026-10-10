"""Extracción común para FastAPI, Streamlit y Drive."""
import base64
import os
from pathlib import Path
import fitz
from docx import Document
from pypdf import PdfReader
from configuracion import DATA_DIR
from openai import OpenAI
MODELO_OCR = os.getenv('GRIFO_MODELO_OCR', 'gpt-5.6-luna')
client = OpenAI()

def extraer_txt(
    ruta: Path
):

    try:

        return ruta.read_text(
            encoding="utf-8"
        )

    except UnicodeDecodeError:

        return ruta.read_text(
            encoding="latin-1"
        )

def pagina_pdf_a_data_url(
    pagina,
    zoom: float = 2.0,
):

    matriz = fitz.Matrix(
        zoom,
        zoom,
    )

    pixmap = pagina.get_pixmap(
        matrix=matriz,
        alpha=False,
    )

    png_bytes = pixmap.tobytes(
        "png"
    )

    base64_png = (
        base64
        .b64encode(
            png_bytes
        )
        .decode(
            "utf-8"
        )
    )

    return (
        "data:image/png;base64,"
        +
        base64_png
    )

def ocr_imagen_openai(
    image_data_url: str,
    numero_pagina: int,
):

    print(
        f"[OCR] Página {numero_pagina}"
    )

    respuesta = (
        client.responses.create(
            model=
                MODELO_OCR,

            input=[
                {
                    "role":
                        "user",

                    "content": [
                        {
                            "type":
                                "input_text",

                            "text":
                                (
                                    "Transcribe fielmente todo el texto "
                                    "legible de esta página. "
                                    "No resumas, no expliques y no inventes. "
                                    "Mantén títulos, listas, fechas, números "
                                    "y tablas de manera legible. "
                                    "Devuelve solamente la transcripción."
                                ),
                        },
                        {
                            "type":
                                "input_image",

                            "image_url":
                                image_data_url,

                            "detail":
                                "high",
                        },
                    ],
                }
            ],
        )
    )

    return (
        respuesta.output_text
        or ""
    ).strip()

def extraer_pdf_con_ocr(
    ruta: Path
):

    documento = fitz.open(
        str(ruta)
    )

    textos = []

    try:

        for numero_pagina, pagina in enumerate(
            documento,
            start=1,
        ):

            imagen = (
                pagina_pdf_a_data_url(
                    pagina
                )
            )

            texto = (
                ocr_imagen_openai(
                    imagen,
                    numero_pagina,
                )
            )

            if texto:

                textos.append(
                    (
                        f"[Página {numero_pagina}]\n"
                        f"{texto}"
                    )
                )

    finally:

        documento.close()

    texto_total = (
        "\n\n"
        .join(
            textos
        )
        .strip()
    )

    if not texto_total:

        raise ValueError(
            "OCR completado pero no encontró texto."
        )

    return texto_total

def extraer_pdf(ruta: Path):
    textos = []
    try:
        documento = fitz.open(str(ruta))
    except Exception:
        # Conservar el lector alternativo para PDF que PyMuPDF no pueda abrir.
        lector = PdfReader(str(ruta))
        if len(lector.pages) > 100:
            raise ValueError('El PDF supera el límite de 100 páginas por carga.')
        for numero, pagina in enumerate(lector.pages, start=1):
            texto = (pagina.extract_text() or '').strip()
            if not texto:
                raise ValueError('Una página requiere OCR, pero el PDF no se pudo renderizar.')
            textos.append(f'[Página {numero}]\n{texto}')
    else:
        with documento:
            if len(documento) > 100:
                raise ValueError('El PDF supera el límite de 100 páginas por carga.')
            for numero, pagina in enumerate(documento, start=1):
                texto = (pagina.get_text('text') or '').strip()
                if not texto:
                    texto = ocr_imagen_openai(pagina_pdf_a_data_url(pagina), numero)
                if texto:
                    textos.append(f'[Página {numero}]\n{texto}')
    if not textos:
        raise ValueError('El PDF no contiene texto legible.')
    return '\n\n'.join(textos)


def extraer_docx(
    ruta: Path
):

    documento = Document(
        str(ruta)
    )

    textos = []

    for parrafo in documento.paragraphs:

        texto = (
            parrafo.text
            or ""
        ).strip()

        if texto:

            textos.append(
                texto
            )

    for tabla in documento.tables:

        for fila in tabla.rows:

            celdas = []

            for celda in fila.cells:

                contenido = (
                    celda.text
                    or ""
                ).strip()

                if contenido:

                    celdas.append(
                        contenido
                    )

            if celdas:

                textos.append(
                    " | ".join(
                        celdas
                    )
                )

    return "\n".join(
        textos
    )

def extraer_texto_documento(
    ruta: Path
):

    extension = (
        ruta.suffix
        .lower()
    )

    if extension == ".txt":
        return extraer_txt(
            ruta
        )

    if extension == ".pdf":
        return extraer_pdf(
            ruta
        )

    if extension == ".docx":
        return extraer_docx(
            ruta
        )

    raise ValueError(
        f"Extensión no soportada: {extension}"
    )
