import json
import os
from pathlib import Path
from configuracion import DATA_DIR

from dotenv import load_dotenv
from openai import OpenAI
from pypdf import PdfReader
from docx import Document

from drive_tools import descargar_archivo_drive


# .env se carga desde configuracion antes de crear clientes.

if not os.getenv("OPENAI_API_KEY"):
    raise RuntimeError(
        "No se encontró OPENAI_API_KEY."
    )

client = OpenAI()

ARCHIVO_INDICE = DATA_DIR / "indice.json"

MODELO_EMBEDDING = "text-embedding-3-small"

TAMANO_FRAGMENTO = 1200
SOLAPAMIENTO = 200


def extraer_txt(ruta: Path) -> str:
    return ruta.read_text(
        encoding="utf-8",
        errors="replace"
    )


def extraer_pdf(ruta: Path) -> str:
    reader = PdfReader(str(ruta))

    paginas = []

    for numero, pagina in enumerate(
        reader.pages,
        start=1
    ):
        texto = pagina.extract_text() or ""

        paginas.append(
            f"\n--- PÁGINA {numero} ---\n{texto}"
        )

    return "\n".join(paginas)


def extraer_docx(ruta: Path) -> str:
    documento = Document(str(ruta))

    bloques = []

    for parrafo in documento.paragraphs:
        texto = parrafo.text.strip()

        if texto:
            bloques.append(texto)

    for numero_tabla, tabla in enumerate(
        documento.tables,
        start=1
    ):
        bloques.append(
            f"\n--- TABLA {numero_tabla} ---"
        )

        for fila in tabla.rows:
            valores = [
                celda.text.strip()
                for celda in fila.cells
            ]

            bloques.append(
                " | ".join(valores)
            )

    return "\n".join(bloques)


def extraer_documento(ruta: Path) -> str:
    extension = ruta.suffix.lower()

    if extension == ".txt":
        return extraer_txt(ruta)

    if extension == ".pdf":
        return extraer_pdf(ruta)

    if extension == ".docx":
        return extraer_docx(ruta)

    raise ValueError(
        f"Formato no soportado para RAG: {extension}"
    )


def dividir_texto(texto: str):
    fragmentos = []

    inicio = 0

    while inicio < len(texto):
        fin = inicio + TAMANO_FRAGMENTO

        fragmento = texto[inicio:fin].strip()

        if fragmento:
            fragmentos.append(fragmento)

        inicio += TAMANO_FRAGMENTO - SOLAPAMIENTO

    return fragmentos


def crear_embedding(texto: str):
    respuesta = client.embeddings.create(
        model=MODELO_EMBEDDING,
        input=texto
    )

    return respuesta.data[0].embedding


def cargar_indice():
    if not ARCHIVO_INDICE.exists():
        return []

    with ARCHIVO_INDICE.open(
        "r",
        encoding="utf-8"
    ) as archivo:
        return json.load(archivo)


def guardar_indice(indice):
    with ARCHIVO_INDICE.open(
        "w",
        encoding="utf-8"
    ) as archivo:
        json.dump(
            indice,
            archivo,
            ensure_ascii=False
        )


def ya_indexado(indice, file_id):
    for item in indice:
        if item.get("drive_file_id") == file_id:
            return True

    return False


def indexar_archivo_drive(file_id: str) -> str:
    print(
        f"[RAG DRIVE] Descargando archivo -> {file_id}"
    )

    indice = cargar_indice()

    if ya_indexado(indice, file_id):
        return (
            "El archivo ya fue indexado anteriormente."
        )

    try:
        ruta = descargar_archivo_drive(
            file_id
        )
    except Exception as error:
        raise RuntimeError("No fue posible descargar el archivo de Drive.") from error

    ruta = Path(ruta)

    try:
        texto = extraer_documento(
            ruta
        )
    except Exception as error:
        raise RuntimeError("No fue posible extraer texto del archivo de Drive.") from error

    if not texto.strip():
        raise ValueError("El documento no contiene texto extraíble.")

    fragmentos = dividir_texto(
        texto
    )

    print(
        f"[RAG DRIVE] Fragmentos -> {len(fragmentos)}"
    )

    for numero, fragmento in enumerate(
        fragmentos,
        start=1
    ):
        print(
            f"[RAG DRIVE] Embedding "
            f"{numero}/{len(fragmentos)}"
        )

        embedding = crear_embedding(
            fragmento
        )

        indice.append({
            "archivo": ruta.name,
            "fragmento": numero,
            "texto": fragmento,
            "embedding": embedding,
            "origen": "google_drive",
            "drive_file_id": file_id,
        })

    guardar_indice(
        indice
    )

    return (
        f"Archivo indexado correctamente.\n"
        f"Nombre: {ruta.name}\n"
        f"Fragmentos agregados: {len(fragmentos)}"
    )