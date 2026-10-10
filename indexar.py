import almacenamiento
from almacenamiento import indice_transaccion
from configuracion import DATA_DIR
import json
import os
from pathlib import Path

import numpy as np
from dotenv import load_dotenv
from openai import OpenAI
from pypdf import PdfReader
from docx import Document


load_dotenv()

if not os.getenv("OPENAI_API_KEY"):
    raise RuntimeError(
        "No se encontró OPENAI_API_KEY. Revisa el archivo .env."
    )

client = OpenAI()

CARPETA_DOCUMENTOS = (DATA_DIR / "documentos")
ARCHIVO_INDICE = (DATA_DIR / "indice.json")

EXTENSIONES = {".txt", ".pdf", ".docx"}

MODELO_EMBEDDING = "text-embedding-3-small"

TAMANO_FRAGMENTO = 1200
SOLAPAMIENTO = 200


def extraer_txt(ruta: Path) -> str:
    return ruta.read_text(
        encoding="utf-8",
        errors="replace"
    )


def extraer_pdf(ruta: Path) -> str:
    reader = PdfReader(str(ruta))

    paginas = []

    for numero, pagina in enumerate(reader.pages, start=1):
        texto = pagina.extract_text() or ""

        paginas.append(
            f"\n--- PÁGINA {numero} ---\n{texto}"
        )

    return "\n".join(paginas)


def extraer_docx(ruta: Path) -> str:
    documento = Document(str(ruta))

    bloques = []

    for parrafo in documento.paragraphs:
        texto = parrafo.text.strip()

        if texto:
            bloques.append(texto)

    for numero_tabla, tabla in enumerate(
        documento.tables,
        start=1
    ):
        bloques.append(
            f"\n--- TABLA {numero_tabla} ---"
        )

        for fila in tabla.rows:
            valores = [
                celda.text.strip()
                for celda in fila.cells
            ]

            bloques.append(
                " | ".join(valores)
            )

    return "\n".join(bloques)


def extraer_documento(ruta):
    from documentos import extraer_texto_documento
    return extraer_texto_documento(ruta)


def dividir_texto(texto: str):
    fragmentos = []

    inicio = 0

    while inicio < len(texto):
        fin = inicio + TAMANO_FRAGMENTO

        fragmento = texto[inicio:fin].strip()

        if fragmento:
            fragmentos.append(fragmento)

        inicio += TAMANO_FRAGMENTO - SOLAPAMIENTO

    return fragmentos


def crear_embedding(texto: str):
    respuesta = client.embeddings.create(
        model=MODELO_EMBEDDING,
        input=texto
    )

    return respuesta.data[0].embedding


@indice_transaccion
def main():
    if not CARPETA_DOCUMENTOS.exists():
        print("La carpeta documentos no existe.")
        return

    archivos = sorted(
        archivo
        for archivo in CARPETA_DOCUMENTOS.iterdir()
        if archivo.is_file()
        and archivo.suffix.lower() in EXTENSIONES
    )

    if not archivos:
        print("No hay documentos compatibles.")
        return

    indice = almacenamiento.cargar_indice()

    print()
    print("=" * 60)
    print("INDEXACIÓN SEMÁNTICA")
    print("=" * 60)
    print()

    for archivo in archivos:
        print(f"Leyendo: {archivo.name}")

        try:
            texto = extraer_documento(archivo)
        except Exception as error:
            print(f"  Error: {error}")
            continue

        if not texto.strip():
            print("  Sin texto extraíble.")
            continue

        fragmentos = dividir_texto(texto)

        print(
            f"  Fragmentos: {len(fragmentos)}"
        )

        indice = [x for x in indice if not (x.get("archivo") == archivo.name and x.get("origen", "local") == "local" and not x.get("drive_file_id"))]

        for numero, fragmento in enumerate(
            fragmentos,
            start=1
        ):
            print(
                f"  Embedding {numero}/"
                f"{len(fragmentos)}"
            )

            embedding = crear_embedding(fragmento)

            indice.append({
                "archivo": archivo.name,
                "fragmento": numero,
                "texto": fragmento,
                "embedding": embedding,
            })

    almacenamiento.guardar_indice(indice)

    print()
    print("=" * 60)
    print("INDEXACIÓN COMPLETADA")
    print("=" * 60)
    print(
        f"Fragmentos indexados: {len(indice)}"
    )
    print(
        f"Archivo generado: {ARCHIVO_INDICE}"
    )


if __name__ == "__main__":
    main()